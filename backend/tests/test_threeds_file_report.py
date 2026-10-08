import os
import tempfile
import unittest
from pathlib import Path

from app.threeds.parsers.models import ThreeDsMessageRow
from app.threeds.services import file_report
from app.threeds.services.csv_storage import resolve_csv_paths_for_dates, save_daily_csvs


def _areq(txn: str, when: str, *, acct: str = "4111111111111111") -> ThreeDsMessageRow:
    return ThreeDsMessageRow(
        log_file="elastic.log",
        message_datetime=when,
        message_type="AReq",
        three_ds_server_trans_id=txn,
        acct_number=acct,
        merchant_name="Test Merchant",
        browser_user_agent="Mozilla/5.0 (Linux; Android 14; SM-S921B) AppleWebKit/537.36",
        purchase_amount="10000",
        purchase_currency="980",
    )


def _ares(txn: str, when: str, *, status: str = "C", reason: str = "") -> ThreeDsMessageRow:
    return ThreeDsMessageRow(
        log_file="elastic.log",
        message_datetime=when,
        message_type="ARes",
        three_ds_server_trans_id=txn,
        trans_status=status,
        trans_status_reason=reason,
    )


def _cres(txn: str, when: str, *, status: str = "Y") -> ThreeDsMessageRow:
    return ThreeDsMessageRow(
        log_file="elastic.log",
        message_datetime=when,
        message_type="CRes",
        three_ds_server_trans_id=txn,
        trans_status=status,
    )


def _rreq(txn: str, when: str, *, status: str = "Y", reason: str = "") -> ThreeDsMessageRow:
    return ThreeDsMessageRow(
        log_file="elastic.log",
        message_datetime=when,
        message_type="RReq",
        three_ds_server_trans_id=txn,
        trans_status=status,
        trans_status_reason=reason,
    )


def _rreq_result(txn: str, when: str, *, status: str, reason: str) -> ThreeDsMessageRow:
    return ThreeDsMessageRow(
        log_file="elastic.log",
        message_datetime=when,
        message_type="RReqResult",
        three_ds_server_trans_id=txn,
        rreq_result_status=status,
        rreq_result_reason=reason,
    )


def _method_expired(txn: str, when: str) -> ThreeDsMessageRow:
    return ThreeDsMessageRow(
        log_file="elastic.log",
        message_datetime=when,
        message_type="MethodExpired",
        three_ds_server_trans_id=txn,
    )


class ThreeDsFileReportTest(unittest.TestCase):
    def test_pivot_row_fields_order(self) -> None:
        self.assertEqual(
            file_report.PIVOT_ROW_FIELDS,
            [
                "general_success",
                "areq_messagedate",
                "browser_os",
                "browser_model",
                "card_scheme",
                "final_cres_status",
                "ares_status",
                "rreq_status",
                "txn_timeline",
                "merchant_name",
                "threedsservertransid",
                "method_expired",
            ],
        )

    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self._saved = {
            key: os.environ.get(key)
            for key in ("THREEDS_CSV_STORAGE_DIR", "THREEDS_REPORT_OUTPUT_DIR", "ELASTIC_TIME_ZONE")
        }
        os.environ["THREEDS_CSV_STORAGE_DIR"] = str(Path(self._tmp.name) / "csv_3ds")
        os.environ["THREEDS_REPORT_OUTPUT_DIR"] = str(Path(self._tmp.name) / "reports_3ds")

    def tearDown(self) -> None:
        file_report.clear_report_cache()
        for key, value in self._saved.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value
        self._tmp.cleanup()

    def _seed_day(self, day: str, txn: str = "T1") -> None:
        save_daily_csvs(
            [
                _areq(txn, f"{day} 10:00:00.000"),
                _ares(txn, f"{day} 10:00:01.000"),
                _cres(txn, f"{day} 10:05:00.000"),
            ]
        )

    def test_default_report_single_day(self) -> None:
        self._seed_day("2026-06-20", "T1")
        result = file_report.run_report_query(
            mode="date",
            date_from="2026-06-20 00:00:00",
            date_to="2026-06-20 23:59:59",
        )
        self.assertEqual(result.row_count, 1)
        self.assertIn("threedsservertransid", result.columns)
        row = result.rows[0]
        self.assertEqual(row["threedsservertransid"], "T1")
        self.assertIn("AReq", row["txn_timeline"])
        self.assertIn("ARes(C+NULL)", row["txn_timeline"])
        self.assertIn("CRes(Y)", row["txn_timeline"])
        self.assertEqual(row["browser_os"], "Android")
        self.assertEqual(row["browser_model"], "Samsung SM-S921B")
        self.assertEqual(row["card_scheme"], "Visa")
        self.assertEqual(row["purchase_amount"], "10000")

    def test_general_success_yes_for_rreq_y(self) -> None:
        save_daily_csvs(
            [
                _areq("T1", "2026-06-20 10:00:00.000"),
                _ares("T1", "2026-06-20 10:00:01.000", status="C"),
                _rreq("T1", "2026-06-20 10:01:00.000", status="Y"),
            ]
        )
        result = file_report.run_report_query(
            mode="date",
            date_from="2026-06-20 00:00:00",
            date_to="2026-06-20 23:59:59",
        )
        self.assertEqual(result.rows[0]["general_success"], "YES")
        self.assertEqual(result.rows[0]["txn_result"], "SUCCESS")

    def test_general_success_no_for_rreq_denied(self) -> None:
        save_daily_csvs(
            [
                _areq("T1", "2026-06-20 10:00:00.000"),
                _ares("T1", "2026-06-20 10:00:01.000", status="C"),
                _rreq("T1", "2026-06-20 10:01:00.000", status="N", reason="14"),
                _rreq_result("T1", "2026-06-20 10:01:00.000", status="DENIED", reason="TXN_TIMED_OUT_AT_ACS"),
            ]
        )
        result = file_report.run_report_query(
            mode="date",
            date_from="2026-06-20 00:00:00",
            date_to="2026-06-20 23:59:59",
        )
        row = result.rows[0]
        self.assertEqual(row["general_success"], "NO")
        self.assertEqual(row["txn_result"], "DENIED")
        self.assertEqual(row["rreq_result_description"], "DENIED / TXN_TIMED_OUT_AT_ACS")
        self.assertIn("Result(DENIED/TXN_TIMED_OUT_AT_ACS)", row["txn_timeline"])

    def test_declined_no_challenge_bucket(self) -> None:
        save_daily_csvs(
            [
                _areq("T1", "2026-06-20 10:00:00.000"),
                _ares("T1", "2026-06-20 10:00:01.000", status="N", reason="01"),
            ]
        )
        result = file_report.run_report_query(
            mode="date",
            date_from="2026-06-20 00:00:00",
            date_to="2026-06-20 23:59:59",
        )
        self.assertEqual(result.rows[0]["txn_result"], "DECLINED_NO_CHALLENGE")
        self.assertEqual(result.rows[0]["general_success"], "NO")

    def test_method_expired_flag(self) -> None:
        save_daily_csvs(
            [
                _areq("T1", "2026-06-20 10:00:00.000"),
                _method_expired("T1", "2026-06-20 10:00:30.000"),
            ]
        )
        result = file_report.run_report_query(
            mode="date",
            date_from="2026-06-20 00:00:00",
            date_to="2026-06-20 23:59:59",
        )
        self.assertEqual(result.rows[0]["method_expired"], "YES")
        self.assertIn("3DSMethodExpired", result.rows[0]["txn_timeline"])

    def test_st_louis_date_field_is_date_only_and_can_roll_back_a_day(self) -> None:
        os.environ["ELASTIC_TIME_ZONE"] = "Europe/Athens"
        save_daily_csvs([_areq("T1", "2026-06-20 03:00:00.000")])
        result = file_report.run_report_query(
            mode="date",
            date_from="2026-06-20 00:00:00",
            date_to="2026-06-20 23:59:59",
        )
        self.assertEqual(result.rows[0]["areq_messagedate_stlouis"], "2026-06-19")

    def test_multi_day_report_counts_all(self) -> None:
        self._seed_day("2026-06-20", "T1")
        self._seed_day("2026-06-21", "T2")
        result = file_report.run_report_query(
            mode="date",
            date_from="2026-06-20 00:00:00",
            date_to="2026-06-21 23:59:59",
        )
        self.assertEqual(result.row_count, 2)
        txns = {row["threedsservertransid"] for row in result.rows}
        self.assertEqual(txns, {"T1", "T2"})

    def test_missing_csv_day_in_range_blocks(self) -> None:
        self._seed_day("2026-06-20", "T1")
        self._seed_day("2026-06-22", "T2")
        with self.assertRaises(ValueError) as ctx:
            file_report.run_report_query(
                mode="date",
                date_from="2026-06-20 00:00:00",
                date_to="2026-06-22 23:59:59",
            )
        self.assertIn("2026-06-21", str(ctx.exception))

    def test_txn_mode_returns_only_requested(self) -> None:
        save_daily_csvs(
            [
                _areq("T1", "2026-06-20 10:00:00.000"),
                _ares("T1", "2026-06-20 10:00:01.000"),
                _areq("T2", "2026-06-20 11:00:00.000"),
                _ares("T2", "2026-06-20 11:00:01.000"),
            ]
        )
        result = file_report.run_report_query(
            mode="txnId",
            date_from="2026-06-20 00:00:00",
            date_to="2026-06-20 23:59:59",
            txn_id="T2",
        )
        self.assertEqual(result.row_count, 1)
        self.assertEqual(result.rows[0]["threedsservertransid"], "T2")

    def test_export_has_data_and_summary_sheets(self) -> None:
        from openpyxl import load_workbook

        save_daily_csvs(
            [
                _areq("T1", "2026-06-20 10:00:00.000"),
                _ares("T1", "2026-06-20 10:00:01.000"),
                _cres("T1", "2026-06-20 10:05:00.000", status="Y"),
                _areq("T2", "2026-06-20 11:00:00.000"),
                _ares("T2", "2026-06-20 11:00:01.000"),
            ]
        )
        export = file_report.export_report_xlsx(
            mode="date",
            date_from="2026-06-20 00:00:00",
            date_to="2026-06-20 23:59:59",
            native_pivot=False,
        )
        workbook = load_workbook(export.output_path)
        self.assertIn("Data", workbook.sheetnames)
        self.assertIn("Summary", workbook.sheetnames)
        self.assertNotIn("Champions", workbook.sheetnames)
        summary = workbook["Summary"]
        header = [cell.value for cell in summary[1]]
        self.assertEqual(header, ["txn_result", "count", "percent"])
        total = sum(row[1].value for row in summary.iter_rows(min_row=2))
        self.assertEqual(total, export.row_count)

    def test_pagination_is_consistent_across_pages(self) -> None:
        save_daily_csvs(
            [
                _areq("T1", "2026-06-20 10:00:00.000"),
                _areq("T2", "2026-06-20 10:01:00.000"),
                _areq("T3", "2026-06-20 10:02:00.000"),
            ]
        )
        kwargs = dict(
            mode="date",
            date_from="2026-06-20 00:00:00",
            date_to="2026-06-20 23:59:59",
        )
        page1 = file_report.run_report_query(limit=2, offset=0, **kwargs)
        page2 = file_report.run_report_query(limit=2, offset=2, **kwargs)
        self.assertEqual(page1.row_count, 3)
        self.assertEqual(page2.row_count, 3)
        self.assertEqual(len(page1.rows), 2)
        self.assertEqual(len(page2.rows), 1)
        seen = {row["threedsservertransid"] for row in page1.rows + page2.rows}
        self.assertEqual(seen, {"T1", "T2", "T3"})

    def test_cache_invalidates_when_data_changes(self) -> None:
        kwargs = dict(
            mode="date",
            date_from="2026-06-20 00:00:00",
            date_to="2026-06-20 23:59:59",
        )
        self._seed_day("2026-06-20", "T1")
        first = file_report.run_report_query(**kwargs)
        self.assertEqual(first.row_count, 1)
        save_daily_csvs(
            [
                _areq("T1", "2026-06-20 10:00:00.000"),
                _areq("T9", "2026-06-20 12:00:00.000"),
            ]
        )
        second = file_report.run_report_query(**kwargs)
        self.assertEqual(second.row_count, 2)

    def test_native_pivot_reports_when_excel_unavailable(self) -> None:
        from unittest.mock import patch

        save_daily_csvs([_areq("T1", "2026-06-20 10:00:00.000")])
        with patch("app.common.excel_pivot.native_pivot_available", return_value=False):
            export = file_report.export_report_xlsx(
                mode="date",
                date_from="2026-06-20 00:00:00",
                date_to="2026-06-20 23:59:59",
                native_pivot=True,
            )
        self.assertFalse(export.pivot_added)
        self.assertIn("Excel is not available", export.pivot_error)

    def test_resolve_missing_day_lists_all_missing(self) -> None:
        self._seed_day("2026-06-20", "T1")
        with self.assertRaises(ValueError) as ctx:
            resolve_csv_paths_for_dates(["2026-06-20", "2026-06-21", "2026-06-22"])
        message = str(ctx.exception)
        self.assertIn("2026-06-21", message)
        self.assertIn("2026-06-22", message)


if __name__ == "__main__":
    unittest.main()
